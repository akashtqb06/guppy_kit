import base64
from io import BytesIO

from docx import Document
from pydantic import BaseModel, Field

from guppy.core.types import ArtifactType, ToolCategory
from guppy.tools.base import BaseTool


class Input(BaseModel):
    docx_base64: str = Field(description="Base64-encoded .docx file content")
    include_headings: bool = Field(
        default=True, description="Mark headings with # prefix based on heading level"
    )
    include_tables: bool = Field(
        default=True, description="Include table content as tab-separated values"
    )


class Output(BaseModel):
    text: str
    paragraph_count: int
    table_count: int
    word_count: int
    heading_count: int


class DocxToTextTool(BaseTool):
    name = "docx-to-text"
    version = "1.0.0"
    category = ToolCategory.DOCUMENTS
    icon = "📝"
    description = "Extract plain text from a Microsoft Word (.docx) file."
    tags = ["docx", "word", "text", "extract", "documents"]  # noqa: RUF012
    input_artifact_types = [ArtifactType.DOCX]  # noqa: RUF012
    output_artifact_type = ArtifactType.TEXT

    input_schema = Input
    output_schema = Output
    config_schema = BaseModel

    async def execute(self, input: Input, config: BaseModel) -> Output:
        docx_bytes = base64.b64decode(input.docx_base64)
        bio = BytesIO(docx_bytes)
        docx = Document(bio)

        extracted_text = []
        heading_count = 0
        paragraph_count = len(docx.paragraphs)

        for para in docx.paragraphs:
            if input.include_headings and para.style.name.startswith("Heading"):
                heading_count += 1
                try:
                    level = int(para.style.name.split()[-1])
                except ValueError:
                    level = 1
                prefix = "#" * level + " "
                extracted_text.append(prefix + para.text)
            else:
                extracted_text.append(para.text)

        table_count = len(docx.tables)
        if input.include_tables:
            for table in docx.tables:
                for row in table.rows:
                    row_data = []
                    for cell in row.cells:
                        row_data.append(cell.text.replace("\n", " ").strip())
                    extracted_text.append("\t".join(row_data))

        full_text = "\n".join(extracted_text)

        return Output(
            text=full_text,
            paragraph_count=paragraph_count,
            table_count=table_count,
            word_count=len(full_text.split()),
            heading_count=heading_count,
        )
