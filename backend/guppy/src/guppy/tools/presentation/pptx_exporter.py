import base64
from io import BytesIO
from typing import Literal

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from pydantic import BaseModel, Field

from guppy.core.types import ArtifactType, ToolCategory
from guppy.tools.base import BaseTool


class SlideInput(BaseModel):
    title: str
    content: str
    notes: str = ""


class Input(BaseModel):
    title: str = Field(description="Presentation title")
    slides: list[SlideInput] = Field(description="List of slides", min_length=1, max_length=50)
    theme: Literal["default", "dark", "minimal"] = Field(default="default")


class Output(BaseModel):
    pptx_base64: str
    slide_count: int
    title: str


class PptxExporterTool(BaseTool):
    name = "pptx-exporter"
    version = "1.0.0"
    category = ToolCategory.PRESENTATION
    icon = "📊"
    description = "Export a slide presentation definition to PowerPoint (.pptx) format."
    tags = ["pptx", "powerpoint", "presentation", "export"]  # noqa: RUF012
    input_artifact_types = [ArtifactType.JSON]  # noqa: RUF012
    output_artifact_type = ArtifactType.PPTX

    input_schema = Input
    output_schema = Output
    config_schema = BaseModel

    async def execute(self, input: Input, config: BaseModel) -> Output:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        theme_colors = {
            "default": {
                "bg": RGBColor(255, 255, 255),
                "title": RGBColor(0x4F, 0x46, 0xE5),
                "text": RGBColor(0, 0, 0),
            },
            "dark": {
                "bg": RGBColor(0x1E, 0x1E, 0x2E),
                "title": RGBColor(0x89, 0xB4, 0xFA),
                "text": RGBColor(255, 255, 255),
            },
            "minimal": {
                "bg": RGBColor(0xFA, 0xFA, 0xFA),
                "title": RGBColor(0x1A, 0x1A, 0x1A),
                "text": RGBColor(0x33, 0x33, 0x33),
            },
        }
        colors = theme_colors.get(input.theme, theme_colors["default"])

        def apply_theme(slide, is_title_slide=False):
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = colors["bg"]

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if shape == slide.shapes.title:
                            run.font.color.rgb = colors["title"]
                        else:
                            run.font.color.rgb = colors["text"]

        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = input.title
        subtitle.text = ""
        apply_theme(slide, is_title_slide=True)

        bullet_slide_layout = prs.slide_layouts[1]
        for slide_data in input.slides:
            slide = prs.slides.add_slide(bullet_slide_layout)
            shapes = slide.shapes
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]

            title_shape.text = slide_data.title
            tf = body_shape.text_frame
            tf.text = ""

            bullets = [b for b in slide_data.content.split("\n") if b.strip()]
            for _i, bullet in enumerate(bullets):
                p = tf.add_paragraph()
                p.text = bullet
                p.level = 0

            if slide_data.notes:
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = slide_data.notes

            apply_theme(slide)

        out_stream = BytesIO()
        prs.save(out_stream)
        out_bytes = out_stream.getvalue()

        return Output(
            pptx_base64=base64.b64encode(out_bytes).decode("utf-8"),
            slide_count=len(prs.slides),
            title=input.title,
        )
