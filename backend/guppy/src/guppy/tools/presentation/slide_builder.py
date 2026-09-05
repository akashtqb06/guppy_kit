"""Slide Builder — generates a real PowerPoint (.pptx) presentation."""

from __future__ import annotations

import base64
from io import BytesIO
from typing import Literal

from pydantic import BaseModel, Field

from guppy.core.types import ArtifactType, ToolCategory
from guppy.tools.base import BaseTool, NoConfig


class SlideContent(BaseModel):
    title: str = Field(description="Slide title")
    content: str = Field(default="", description="Slide body text (bullet points, one per line)")
    notes: str = Field(default="", description="Speaker notes")
    layout: Literal["title", "content", "blank", "two_column"] = Field(default="content")


class SlideBuilderInput(BaseModel):
    title: str = Field(description="Presentation title")
    author: str = Field(default="", description="Author name")
    slides: list[SlideContent] = Field(description="List of slides", min_length=1)
    theme: Literal["default", "dark", "minimal", "corporate", "vibrant"] = Field(default="default")


class SlideBuilderOutput(BaseModel):
    pptx_base64: str = Field(description="Base64-encoded .pptx file")
    slide_count: int
    title: str
    outline: list[str]
    file_size_bytes: int


# Theme colors: (bg_rgb, title_rgb, body_rgb, accent_rgb)
THEMES = {
    "default": {
        "bg": (0xFF, 0xFF, 0xFF),
        "title": (0x1A, 0x1A, 0x2E),
        "body": (0x33, 0x33, 0x33),
        "accent": (0x2B, 0x6C, 0xD4),
    },
    "dark": {
        "bg": (0x1A, 0x1A, 0x2E),
        "title": (0xFF, 0xFF, 0xFF),
        "body": (0xCC, 0xCC, 0xCC),
        "accent": (0x6C, 0x9B, 0xFF),
    },
    "minimal": {
        "bg": (0xFA, 0xFA, 0xFA),
        "title": (0x00, 0x00, 0x00),
        "body": (0x44, 0x44, 0x44),
        "accent": (0x88, 0x88, 0x88),
    },
    "corporate": {
        "bg": (0xF4, 0xF6, 0xF8),
        "title": (0x0D, 0x2B, 0x55),
        "body": (0x1C, 0x3A, 0x6E),
        "accent": (0x00, 0x82, 0xCA),
    },
    "vibrant": {
        "bg": (0xFE, 0xFE, 0xFE),
        "title": (0x6B, 0x21, 0xA8),
        "body": (0x1E, 0x1E, 0x2E),
        "accent": (0xEC, 0x48, 0x99),
    },
}


class SlideBuilderTool(BaseTool[SlideBuilderInput, NoConfig, SlideBuilderOutput]):
    name = "slide-builder"
    version = "2.0.0"
    category = ToolCategory.PRESENTATION
    description = "Build a real PowerPoint (.pptx) presentation with themed slides."
    tags = ["presentation", "powerpoint", "pptx", "slides"]  # noqa: RUF012
    output_artifact_type = ArtifactType.PPTX
    input_artifact_types = []  # noqa: RUF012
    icon = "presentation"
    input_schema = SlideBuilderInput
    output_schema = SlideBuilderOutput
    config_schema = NoConfig

    async def execute(self, input: SlideBuilderInput, config: NoConfig) -> SlideBuilderOutput:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        prs = Presentation()
        # Widescreen 16:9
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        theme = THEMES.get(input.theme, THEMES["default"])
        bg_color = RGBColor(*theme["bg"])
        title_color = RGBColor(*theme["title"])
        body_color = RGBColor(*theme["body"])
        accent_color = RGBColor(*theme["accent"])

        blank_layout = prs.slide_layouts[6]  # Blank

        def set_bg(slide):
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = bg_color

        def add_text_box(
            slide,
            text,
            left,
            top,
            width,
            height,
            font_size,
            bold=False,
            color=None,
            align=PP_ALIGN.LEFT,
        ):
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = align
            run = p.add_run()
            run.text = text
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color or title_color
            return txBox

        def add_accent_bar(slide):
            """Thin accent bar at bottom"""
            left = Inches(0)
            top = prs.slide_height - Inches(0.08)
            width = prs.slide_width
            height = Inches(0.08)
            shape = slide.shapes.add_shape(1, left, top, width, height)  # 1=RECTANGLE
            shape.fill.solid()
            shape.fill.fore_color.rgb = accent_color
            shape.line.fill.background()

        for i, slide_data in enumerate(input.slides):
            slide = prs.slides.add_slide(blank_layout)
            set_bg(slide)
            add_accent_bar(slide)

            if (i == 0 and slide_data.layout == "title") or (i == 0 and len(input.slides) > 1):
                # Title slide
                add_text_box(
                    slide,
                    input.title,
                    Inches(1.5),
                    Inches(2.5),
                    Inches(10),
                    Inches(1.5),
                    font_size=40,
                    bold=True,
                    color=title_color,
                    align=PP_ALIGN.CENTER,
                )
                if slide_data.content:
                    add_text_box(
                        slide,
                        slide_data.content,
                        Inches(2),
                        Inches(4.2),
                        Inches(9),
                        Inches(1),
                        font_size=20,
                        bold=False,
                        color=body_color,
                        align=PP_ALIGN.CENTER,
                    )
                if input.author:
                    add_text_box(
                        slide,
                        input.author,
                        Inches(2),
                        Inches(5.5),
                        Inches(9),
                        Inches(0.5),
                        font_size=14,
                        bold=False,
                        color=RGBColor(*theme["accent"]),
                        align=PP_ALIGN.CENTER,
                    )
            else:
                # Content slide
                # Title
                add_text_box(
                    slide,
                    slide_data.title,
                    Inches(0.5),
                    Inches(0.3),
                    Inches(12.3),
                    Inches(1.0),
                    font_size=28,
                    bold=True,
                    color=title_color,
                )
                # Divider line
                line = slide.shapes.add_shape(
                    1, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.02)
                )
                line.fill.solid()
                line.fill.fore_color.rgb = accent_color
                line.line.fill.background()

                # Body content (bullets)
                if slide_data.content:
                    txBox = slide.shapes.add_textbox(
                        Inches(0.5), Inches(1.6), Inches(12.3), Inches(5.2)
                    )
                    tf = txBox.text_frame
                    tf.word_wrap = True
                    lines = slide_data.content.strip().split("\n")
                    for j, line_text in enumerate(lines):
                        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                        p.space_before = Pt(4)
                        run = p.add_run()
                        bullet_text = line_text.lstrip("- •*").strip()
                        run.text = f"  • {bullet_text}" if bullet_text else ""
                        run.font.size = Pt(18)
                        run.font.color.rgb = body_color

        buf = BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        return SlideBuilderOutput(
            pptx_base64=base64.b64encode(pptx_bytes).decode(),
            slide_count=len(input.slides),
            title=input.title,
            outline=[s.title for s in input.slides],
            file_size_bytes=len(pptx_bytes),
        )
